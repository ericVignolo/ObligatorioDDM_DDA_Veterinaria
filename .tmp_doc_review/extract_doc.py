from __future__ import annotations

import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(path)

    paragraphs = []
    run_formats = Counter()
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                run_formats[
                    (
                        run.font.name or "(inherit)",
                        round(run.font.size.pt, 1) if run.font.size else "(inherit)",
                        bool(run.bold),
                        bool(run.italic),
                        str(run.font.color.rgb) if run.font.color and run.font.color.rgb else "(inherit)",
                    )
                ] += len(run.text)
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else None,
                    "text": text,
                }
            )

    tables = []
    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip().replace("\n", " | ") for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    sections = []
    for index, section in enumerate(doc.sections, start=1):
        sections.append(
            {
                "index": index,
                "width_inches": round(section.page_width.inches, 2),
                "height_inches": round(section.page_height.inches, 2),
                "margins_inches": {
                    "top": round(section.top_margin.inches, 2),
                    "bottom": round(section.bottom_margin.inches, 2),
                    "left": round(section.left_margin.inches, 2),
                    "right": round(section.right_margin.inches, 2),
                },
                "header": " || ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip()),
                "footer": " || ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip()),
            }
        )

    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        package = {
            "media": [name for name in names if name.startswith("word/media/")],
            "has_comments": "word/comments.xml" in names,
            "has_footnotes": "word/footnotes.xml" in names,
            "has_endnotes": "word/endnotes.xml" in names,
            "has_toc_field": "TOC" in xml,
            "drawings": xml.count("<w:drawing"),
            "page_breaks": xml.count('w:type="page"'),
        }

    props = doc.core_properties
    named_styles = []
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Paragraph"]:
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        named_styles.append(
            {
                "name": style_name,
                "font": style.font.name,
                "size_pt": round(style.font.size.pt, 1) if style.font.size else None,
                "bold": style.font.bold,
                "italic": style.font.italic,
                "color": str(style.font.color.rgb) if style.font.color and style.font.color.rgb else None,
            }
        )
    report = {
        "file": str(path),
        "core_properties": {
            "title": props.title,
            "subject": props.subject,
            "author": props.author,
            "last_modified_by": props.last_modified_by,
            "created": props.created.isoformat() if props.created else None,
            "modified": props.modified.isoformat() if props.modified else None,
        },
        "counts": {
            "paragraphs_nonempty": len(paragraphs),
            "tables": len(tables),
            "inline_shapes": len(doc.inline_shapes),
            "sections": len(doc.sections),
            "styles": Counter(p["style"] for p in paragraphs),
        },
        "named_styles": named_styles,
        "run_format_char_counts": [
            {
                "font": key[0],
                "size_pt": key[1],
                "bold": key[2],
                "italic": key[3],
                "color": key[4],
                "characters": count,
            }
            for key, count in run_formats.most_common()
        ],
        "sections": sections,
        "package": package,
        "paragraphs": paragraphs,
        "tables": tables,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2, default=dict))


if __name__ == "__main__":
    main()
